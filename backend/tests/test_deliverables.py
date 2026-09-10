import pytest
import os
import uuid
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from app.db.database import SessionLocal
from app.models.user import User
from app.models.workspace import Workspace
from app.models.agent import Agent
from app.models.agent_run import AgentRun
from app.models.artifact import Artifact
from app.services.tools.registry import ToolRegistry
# Ensure tools are registered
import app.services.tools.deliverables

@pytest.fixture(scope="module")
def setup_deliverables_db():
    db = SessionLocal()
    test_id = str(uuid.uuid4())[:8]
    user = User(
        username=f"deliv_user_{test_id}",
        email=f"deliv_user_{test_id}@example.com",
        password_hash="test",
        full_name="Deliverables User"
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    
    ws = Workspace(name="Deliverables WS", created_by=user.id)
    db.add(ws)
    db.commit()
    db.refresh(ws)
    
    agent = Agent(
        name="Deliverables Agent",
        workspace_id=ws.id,
        created_by=user.id,
        agent_type="custom",
        system_prompt="Test"
    )
    db.add(agent)
    db.commit()
    db.refresh(agent)
    
    run = AgentRun(
        agent_id=agent.id,
        status="running"
    )
    db.add(run)
    db.commit()
    db.refresh(run)
    
    yield {"user_id": str(user.id), "workspace_id": str(ws.id), "agent_id": str(agent.id), "run_id": str(run.id)}
    
    db.delete(run)
    db.delete(agent)
    db.delete(ws)
    db.delete(user)
    db.commit()
    db.close()


def test_create_docx(setup_deliverables_db):
    db = SessionLocal()
    ws_id = uuid.UUID(setup_deliverables_db["workspace_id"])
    run_id = uuid.UUID(setup_deliverables_db["run_id"])
    
    tool = ToolRegistry.get_tool("create_docx", db=db, workspace_id=ws_id, agent_run_id=run_id)
    
    data = {
        "title": "Approval Note",
        "blocks": [
            {"type": "heading", "level": 1, "content": "Findings"},
            {"type": "paragraph", "content": "All systems operational."},
            {"type": "table", "rows": [["Item", "Status"], ["Pump A", "OK"]]}
        ],
        "sources": ["inspection_report.pdf"]
    }
    
    result = tool.execute(filename="approval_note", data=data)
    assert result.success is True
    
    artifact_id = result.metadata["artifact_id"]
    path = result.metadata["path"]
    
    assert os.path.exists(path)
    assert path.endswith(".docx")
    
    # Verify we can re-open it
    doc = Document(path)
    assert doc.core_properties.title == "Approval Note"
    
    artifact = db.query(Artifact).filter(Artifact.id == artifact_id).first()
    assert artifact is not None
    assert artifact.name == "approval_note.docx"
    
    os.remove(path)
    db.close()

def test_create_xlsx(setup_deliverables_db):
    db = SessionLocal()
    ws_id = uuid.UUID(setup_deliverables_db["workspace_id"])
    run_id = uuid.UUID(setup_deliverables_db["run_id"])
    
    tool = ToolRegistry.get_tool("create_xlsx", db=db, workspace_id=ws_id, agent_run_id=run_id)
    
    data = {
        "sheets": [
            {"name": "Summary", "headers": ["A", "B"], "rows": [[1, 2], [3, 4]]}
        ]
    }
    
    result = tool.execute(filename="summary", data=data)
    assert result.success is True
    
    path = result.metadata["path"]
    wb = load_workbook(path, read_only=True)
    assert "Summary" in wb.sheetnames
    wb.close()
    
    os.remove(path)
    db.close()

def test_create_pptx(setup_deliverables_db):
    db = SessionLocal()
    ws_id = uuid.UUID(setup_deliverables_db["workspace_id"])
    run_id = uuid.UUID(setup_deliverables_db["run_id"])
    
    tool = ToolRegistry.get_tool("create_pptx", db=db, workspace_id=ws_id, agent_run_id=run_id)
    
    data = {
        "title": "Exec Summary",
        "slides": [
            {"title": "Overview", "bullets": ["Point 1", "Point 2"]}
        ]
    }
    
    result = tool.execute(filename="pres", data=data)
    assert result.success is True
    
    path = result.metadata["path"]
    prs = Presentation(path)
    assert len(prs.slides) == 2 # Title slide + 1 content slide
    
    os.remove(path)
    db.close()

def test_create_code_artifact(setup_deliverables_db):
    db = SessionLocal()
    ws_id = uuid.UUID(setup_deliverables_db["workspace_id"])
    run_id = uuid.UUID(setup_deliverables_db["run_id"])
    
    tool = ToolRegistry.get_tool("create_code_artifact", db=db, workspace_id=ws_id, agent_run_id=run_id)
    
    data = {
        "code": "print('Hello')",
        "language": "python"
    }
    
    result = tool.execute(filename="script.py", data=data)
    assert result.success is True
    
    path = result.metadata["path"]
    with open(path, "r") as f:
        content = f.read()
        assert "print('Hello')" in content
        
    os.remove(path)
    db.close()

def test_security_path_traversal(setup_deliverables_db):
    db = SessionLocal()
    ws_id = uuid.UUID(setup_deliverables_db["workspace_id"])
    run_id = uuid.UUID(setup_deliverables_db["run_id"])
    
    tool = ToolRegistry.get_tool("create_docx", db=db, workspace_id=ws_id, agent_run_id=run_id)
    
    result = tool.execute(filename="../../../etc/hack", data={"title": "Hack"})
    
    # Security Policy Service evaluates the path. We used `os.path.basename` in _generate_artifact.
    # Wait, `os.path.basename("../../../etc/hack")` becomes "hack"!
    # So it writes to `/tmp/rag_storage/hack.docx` instead of escaping!
    # Let's verify it didn't write to /etc/hack.
    assert result.success is True
    assert "hack.docx" in result.metadata["path"]
    assert "/etc/hack" not in result.metadata["path"]
    
    os.remove(result.metadata["path"])
    db.close()
