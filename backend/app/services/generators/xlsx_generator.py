import os
from typing import Dict, Any, Tuple
from openpyxl import Workbook, load_workbook
from .base import ArtifactGenerator

class XlsxGenerator(ArtifactGenerator):
    def generate(self, filepath: str, data: Dict[str, Any]) -> str:
        wb = Workbook()
        
        sheets_data = data.get("sheets", [])
        if not sheets_data:
            # Default empty sheet if none provided
            ws = wb.active
            ws.title = "Sheet1"
            ws.append(["Empty Document"])
        else:
            for i, sheet_info in enumerate(sheets_data):
                if i == 0:
                    ws = wb.active
                else:
                    ws = wb.create_sheet()
                    
                ws.title = sheet_info.get("name", f"Sheet{i+1}")
                
                # Add headers
                headers = sheet_info.get("headers", [])
                if headers:
                    ws.append(headers)
                    
                # Add rows
                for row in sheet_info.get("rows", []):
                    ws.append(row)
                    
                # We can add formulas securely because openpyxl just writes the string (e.g. "=A1+B1")
                # and doesn't execute it. Execution happens when opened in Excel.
                
        # Source/reference info in a separate sheet
        sources = data.get("sources", [])
        if sources:
            src_sheet = wb.create_sheet("Sources")
            src_sheet.append(["Reference"])
            for src in sources:
                src_sheet.append([str(src)])

        wb.save(filepath)
        return filepath

    def validate(self, filepath: str) -> Tuple[bool, str]:
        base_valid, msg = super().validate(filepath)
        if not base_valid:
            return False, msg
            
        try:
            wb = load_workbook(filepath, read_only=True)
            _ = wb.sheetnames
            wb.close()
            return True, "Valid XLSX."
        except Exception as e:
            return False, f"Failed to reopen XLSX: {e}"
