import os
from typing import Dict, Any, Tuple
from docx import Document
from .base import ArtifactGenerator

class DocxGenerator(ArtifactGenerator):
    def generate(self, filepath: str, data: Dict[str, Any]) -> str:
        doc = Document()
        
        # Add basic properties
        core_properties = doc.core_properties
        core_properties.title = data.get("title", "Generated Document")
        core_properties.author = data.get("author", "AI Agent")
        
        # Add content blocks
        if "title" in data:
            doc.add_heading(data["title"], level=0)
            
        for block in data.get("blocks", []):
            b_type = block.get("type")
            b_content = block.get("content")
            
            if not b_content:
                continue
                
            if b_type == "heading":
                level = block.get("level", 1)
                doc.add_heading(b_content, level=level)
            elif b_type == "paragraph":
                doc.add_paragraph(b_content)
            elif b_type == "bullet":
                doc.add_paragraph(b_content, style='List Bullet')
            elif b_type == "table":
                rows = block.get("rows", [])
                if rows:
                    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                    table.style = 'Table Grid'
                    for i, row in enumerate(rows):
                        row_cells = table.rows[i].cells
                        for j, cell_val in enumerate(row):
                            row_cells[j].text = str(cell_val)
                            
        # Source / reference section
        sources = data.get("sources", [])
        if sources:
            doc.add_heading("Sources & References", level=1)
            for src in sources:
                doc.add_paragraph(f"- {src}", style='List Bullet')

        doc.save(filepath)
        return filepath

    def validate(self, filepath: str) -> Tuple[bool, str]:
        base_valid, msg = super().validate(filepath)
        if not base_valid:
            return False, msg
            
        try:
            doc = Document(filepath)
            # Just touch a property to verify it opens as a docx
            _ = len(doc.paragraphs)
            return True, "Valid DOCX."
        except Exception as e:
            return False, f"Failed to reopen DOCX: {e}"
