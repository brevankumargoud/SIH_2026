import os
from typing import Dict, Any, Tuple
from pptx import Presentation
from .base import ArtifactGenerator

class PptxGenerator(ArtifactGenerator):
    def generate(self, filepath: str, data: Dict[str, Any]) -> str:
        prs = Presentation()
        
        # Add Title Slide
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = data.get("title", "Generated Presentation")
        subtitle.text = data.get("subtitle", "AI Agent Deliverable")
        
        # Add content slides
        bullet_slide_layout = prs.slide_layouts[1]
        for slide_data in data.get("slides", []):
            s = prs.slides.add_slide(bullet_slide_layout)
            shapes = s.shapes
            title_shape = shapes.title
            body_shape = shapes.placeholders[1]
            
            title_shape.text = slide_data.get("title", "Slide")
            
            tf = body_shape.text_frame
            for i, bullet in enumerate(slide_data.get("bullets", [])):
                if i == 0:
                    tf.text = bullet
                else:
                    p = tf.add_paragraph()
                    p.text = bullet
                    
        # Source / reference section
        sources = data.get("sources", [])
        if sources:
            s = prs.slides.add_slide(bullet_slide_layout)
            s.shapes.title.text = "Sources & References"
            tf = s.shapes.placeholders[1].text_frame
            for i, src in enumerate(sources):
                if i == 0:
                    tf.text = str(src)
                else:
                    p = tf.add_paragraph()
                    p.text = str(src)

        prs.save(filepath)
        return filepath

    def validate(self, filepath: str) -> Tuple[bool, str]:
        base_valid, msg = super().validate(filepath)
        if not base_valid:
            return False, msg
            
        try:
            prs = Presentation(filepath)
            _ = len(prs.slides)
            return True, "Valid PPTX."
        except Exception as e:
            return False, f"Failed to reopen PPTX: {e}"
